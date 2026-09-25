"""Excel export for the company onboarding table."""

from __future__ import annotations

from datetime import date
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.models.onboarding import OnboardingCell, OnboardingColumn, OnboardingPlan
from app.services.onboarding import onboarding_plan_query, plan_dict


HEADER_FILL = PatternFill(fill_type="solid", fgColor="E8EEF7")
COMPLETED_FILL = PatternFill(fill_type="solid", fgColor="E2F5E9")
COMPLETED_FONT = Font(color="16803C")
DATE_FORMAT = "DD.MM.YYYY"


def _cell_value(column: OnboardingColumn, cell: OnboardingCell | None):
    if cell is None:
        return None
    if cell.is_not_required:
        return "Не нужно"
    if column.field_type == "text":
        return cell.text_value
    if column.field_type == "date":
        return cell.date_value
    if column.field_type == "checkbox":
        return "Выполнено" if cell.is_completed else "Не выполнено"
    if cell.is_completed:
        if cell.completed_date:
            return f"Выполнено: {cell.completed_date.strftime('%d.%m.%Y')}"
        return "Выполнено"
    return cell.planned_date


def build_onboarding_workbook(company_id: int) -> BytesIO:
    columns = (
        OnboardingColumn.query.filter_by(company_id=company_id, is_archived=False)
        .order_by(OnboardingColumn.sort_order.asc(), OnboardingColumn.id.asc())
        .all()
    )
    plans = onboarding_plan_query(company_id).order_by(OnboardingPlan.id.asc()).all()
    serialized_plans = [plan_dict(plan) for plan in plans]
    cells_by_plan = [
        {cell.column_id: cell for cell in plan.cells}
        for plan in plans
    ]

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "План обучения"

    sheet.append(["Этап / сотрудник", *[plan["full_name"] or "—" for plan in serialized_plans]])
    sheet.append(["№", *range(1, len(plans) + 1)])
    sheet.append(["Грейд", *[plan["grade"] or "—" for plan in serialized_plans]])

    for column in columns:
        row = [column.title]
        for cells in cells_by_plan:
            row.append(_cell_value(column, cells.get(column.id)))
        sheet.append(row)
        row_number = sheet.max_row
        for column_number, cells in enumerate(cells_by_plan, start=2):
            cell = cells.get(column.id)
            exported = sheet.cell(row=row_number, column=column_number)
            if isinstance(exported.value, date):
                exported.number_format = DATE_FORMAT
            if cell and cell.is_completed and not cell.is_not_required:
                exported.fill = COMPLETED_FILL
                exported.font = COMPLETED_FONT

    for cell in sheet[1]:
        cell.fill = HEADER_FILL
        cell.font = Font(bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for row_number in range(2, sheet.max_row + 1):
        sheet.cell(row=row_number, column=1).font = Font(bold=True)
    for row in sheet.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    sheet.freeze_panes = "B2"
    sheet.column_dimensions["A"].width = 28
    for column_number in range(2, sheet.max_column + 1):
        sheet.column_dimensions[get_column_letter(column_number)].width = 24

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output

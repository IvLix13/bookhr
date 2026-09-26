"""Reward service."""

from __future__ import annotations

from datetime import date
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font

from app.extensions import db
from app.models import Employment, Reward, RewardStatus
from app.models.reward import REWARD_STATUS_LABELS
from app.services.audit import log_audit
from app.utils.dates import today_moscow


def _parse_date(value: str | None) -> date | None:
    if not value:
        return None
    return date.fromisoformat(value)


def _validate_status(status: str) -> RewardStatus:
    try:
        return RewardStatus(status)
    except ValueError as exc:
        raise ValueError(f"Invalid reward status: {status}") from exc


def create_reward(
    employment_id: int,
    reward_type: str,
    status: str = RewardStatus.NOT_DELIVERED.value,
    directive_text: str | None = None,
    status_changed_date: date | None = None,
    delivered_date: date | None = None,
    notes: str | None = None,
) -> Reward:
    employment = db.session.get(Employment, employment_id)
    if not employment:
        raise ValueError("Employment not found")

    reward_type = reward_type.strip()
    if not reward_type:
        raise ValueError("Reward type is required")

    reward_status = _validate_status(status)
    reward = Reward(
        employment_id=employment_id,
        reward_type=reward_type,
        status=reward_status.value,
        directive_text=directive_text,
        notes=notes,
    )

    parsed_status_changed_date = (
        _parse_date(status_changed_date)
        if isinstance(status_changed_date, str)
        else status_changed_date
    )
    if parsed_status_changed_date:
        reward.status_changed_date = parsed_status_changed_date

    parsed_delivered_date = (
        _parse_date(delivered_date) if isinstance(delivered_date, str) else delivered_date
    )
    if reward_status == RewardStatus.DELIVERED:
        reward.delivered_date = parsed_delivered_date or today_moscow()
    elif parsed_delivered_date:
        reward.delivered_date = parsed_delivered_date

    db.session.add(reward)
    db.session.flush()
    log_audit(
        "create",
        "reward",
        reward.id,
        None,
        {"employment_id": employment_id, "reward_type": reward_type, "status": reward.status},
    )
    return reward


def update_reward(reward: Reward, payload: dict) -> Reward:
    old_value: dict = {
        "status": reward.status,
        "reward_type": reward.reward_type,
        "directive_text": reward.directive_text,
        "status_changed_date": (
            reward.status_changed_date.isoformat() if reward.status_changed_date else None
        ),
        "delivered_date": reward.delivered_date.isoformat() if reward.delivered_date else None,
        "notes": reward.notes,
    }

    if "reward_type" in payload:
        reward_type = payload["reward_type"].strip()
        if not reward_type:
            raise ValueError("Reward type is required")
        reward.reward_type = reward_type

    if "directive_text" in payload:
        reward.directive_text = payload["directive_text"]

    if "notes" in payload:
        reward.notes = payload["notes"]

    if "status_changed_date" in payload:
        reward.status_changed_date = _parse_date(payload["status_changed_date"])

    if "delivered_date" in payload:
        reward.delivered_date = _parse_date(payload["delivered_date"])

    if "status" in payload:
        new_status = _validate_status(payload["status"])
        if new_status.value != reward.status:
            if "status_changed_date" not in payload and reward.status_changed_date is None:
                reward.status_changed_date = today_moscow()
        if new_status == RewardStatus.DELIVERED and reward.delivered_date is None:
            reward.delivered_date = today_moscow()
        reward.status = new_status.value
    elif "delivered_date" not in payload and reward.status == RewardStatus.DELIVERED.value:
        if reward.delivered_date is None:
            reward.delivered_date = today_moscow()

    log_audit(
        "update",
        "reward",
        reward.id,
        old_value,
        {
            "status": reward.status,
            "reward_type": reward.reward_type,
            "directive_text": reward.directive_text,
            "status_changed_date": (
                reward.status_changed_date.isoformat() if reward.status_changed_date else None
            ),
            "delivered_date": reward.delivered_date.isoformat() if reward.delivered_date else None,
            "notes": reward.notes,
        },
    )
    return reward


def list_rewards_for_company(
    company_id: int,
    status: str | None = None,
    employment_id: int | None = None,
) -> list[Reward]:
    query = (
        Reward.query.join(Employment, Reward.employment_id == Employment.id)
        .filter(Employment.company_id == company_id)
        .order_by(Reward.updated_at.desc())
    )
    if status:
        query = query.filter(Reward.status == status)
    if employment_id:
        query = query.filter(Reward.employment_id == employment_id)
    return query.all()


def reward_status_statistics(company_id: int) -> dict:
    """Count every reward in the company, including statuses that are unused."""
    rows = (
        db.session.query(Reward.status, db.func.count(Reward.id))
        .join(Employment, Reward.employment_id == Employment.id)
        .filter(Employment.company_id == company_id)
        .group_by(Reward.status)
        .all()
    )
    counts = {status: count for status, count in rows}
    items = [
        {
            "status": status.value,
            "label": REWARD_STATUS_LABELS[status.value],
            "count": int(counts.get(status.value, 0)),
        }
        for status in RewardStatus
    ]
    known = {status.value for status in RewardStatus}
    for status, count in sorted(rows, key=lambda item: item[0] or ""):
        if status not in known:
            items.append({"status": status, "label": status or "—", "count": int(count)})
    return {"items": items, "total": sum(item["count"] for item in items)}


def build_reward_statistics_workbook(company_id: int) -> BytesIO:
    stats = reward_status_statistics(company_id)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Статистика поощрений"
    sheet.append(["Состояние", "Количество"])
    for item in stats["items"]:
        sheet.append([item["label"], item["count"]])
    sheet.append(["Всего", stats["total"]])
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    sheet.cell(row=sheet.max_row, column=1).font = Font(bold=True)
    sheet.column_dimensions["A"].width = 24
    sheet.column_dimensions["B"].width = 16
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output

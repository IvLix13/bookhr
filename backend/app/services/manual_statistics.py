"""Manual statistics Excel import/export helpers."""

from __future__ import annotations

from datetime import date, datetime, time
from io import BytesIO
from pathlib import Path
import re
from typing import BinaryIO

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils.exceptions import InvalidFileException

from app.extensions import db
from app.models import ManualStatisticsSnapshot

SHEET_TITLE = "Ручная статистика"
LABEL_HEADER = "Наименование параметра"
VALUE_HEADER = "Значение"
BLOCK_MARKER = re.compile(r"^Блок\s+\d+$", re.IGNORECASE)
BLOCK_COLUMNS = ((1, "left", "левой"), (4, "right", "правой"))


def _cell_to_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, (date, datetime, time)):
        return value.isoformat()
    return str(value).strip()


def build_manual_statistics_template() -> BytesIO:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = SHEET_TITLE
    examples = (
        (1, "Договоры", (("Заключено за месяц", 15), ("Заключено за год", 87))),
        (4, "Аттестация", (("Проведено за месяц", 21), ("Проведено за год", 96))),
    )
    for start_column, title, items in examples:
        worksheet.cell(row=1, column=start_column, value="Блок 1")
        worksheet.cell(row=2, column=start_column, value=title)
        worksheet.cell(row=3, column=start_column, value=LABEL_HEADER)
        worksheet.cell(row=3, column=start_column + 1, value=VALUE_HEADER)
        for row_number, (label, value) in enumerate(items, start=4):
            worksheet.cell(row=row_number, column=start_column, value=label)
            worksheet.cell(row=row_number, column=start_column + 1, value=value)

        worksheet.cell(row=1, column=start_column).font = Font(
            bold=True, color="FFFFFF"
        )
        worksheet.cell(row=1, column=start_column).fill = PatternFill(
            "solid", fgColor="2F6FED"
        )
        worksheet.cell(row=2, column=start_column).font = Font(bold=True, size=12)
        for column in (start_column, start_column + 1):
            worksheet.cell(row=3, column=column).font = Font(bold=True)
            worksheet.cell(row=3, column=column).fill = PatternFill(
                "solid", fgColor="DDE8FF"
            )

    worksheet.freeze_panes = "A4"
    worksheet.column_dimensions["A"].width = 38
    worksheet.column_dimensions["B"].width = 20
    worksheet.column_dimensions["C"].width = 4
    worksheet.column_dimensions["D"].width = 38
    worksheet.column_dimensions["E"].width = 20

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


def _parse_block_column(worksheet, start_column: int, column: str, side: str):
    blocks: list[dict[str, object]] = []
    row_number = 1

    while row_number <= worksheet.max_row:
        label = _cell_to_text(worksheet.cell(row=row_number, column=start_column).value)
        value = _cell_to_text(worksheet.cell(row=row_number, column=start_column + 1).value)
        if not label and not value:
            row_number += 1
            continue
        if not BLOCK_MARKER.fullmatch(label):
            raise ValueError(
                f"В {side} колонке, строка {row_number}: ожидается маркер «Блок N»"
            )
        if value:
            raise ValueError(
                f"В {side} колонке, строка {row_number}: рядом с маркером блока должно быть пусто"
            )

        title_row = row_number + 1
        title = _cell_to_text(worksheet.cell(row=title_row, column=start_column).value)
        title_value = _cell_to_text(
            worksheet.cell(row=title_row, column=start_column + 1).value
        )
        if not title:
            raise ValueError(
                f"В {side} колонке, строка {title_row}: не указано название блока"
            )
        if title_value:
            raise ValueError(
                f"В {side} колонке, строка {title_row}: рядом с названием блока должно быть пусто"
            )

        header_row = row_number + 2
        header_label = _cell_to_text(
            worksheet.cell(row=header_row, column=start_column).value
        )
        header_value = _cell_to_text(
            worksheet.cell(row=header_row, column=start_column + 1).value
        )
        if header_label != LABEL_HEADER or header_value != VALUE_HEADER:
            raise ValueError(
                f"В {side} колонке, строка {header_row}: ожидаются заголовки "
                f"«{LABEL_HEADER}» и «{VALUE_HEADER}»"
            )

        items: list[dict[str, str]] = []
        row_number = header_row + 1
        while row_number <= worksheet.max_row:
            item_label = _cell_to_text(
                worksheet.cell(row=row_number, column=start_column).value
            )
            item_value = _cell_to_text(
                worksheet.cell(row=row_number, column=start_column + 1).value
            )
            if BLOCK_MARKER.fullmatch(item_label):
                break
            if not item_label and not item_value:
                row_number += 1
                continue
            if not item_label:
                raise ValueError(
                    f"В {side} колонке, строка {row_number}: не указано наименование параметра"
                )
            items.append({"label": item_label, "value": item_value})
            row_number += 1

        if items:
            blocks.append({"column": column, "title": title, "items": items})

    return blocks


def parse_manual_statistics_workbook(stream: BinaryIO) -> list[dict[str, object]]:
    try:
        workbook = load_workbook(stream, read_only=True, data_only=True)
    except (InvalidFileException, OSError, ValueError) as exc:
        raise ValueError("Не удалось прочитать Excel-файл") from exc

    try:
        worksheet = workbook.active
        blocks: list[dict[str, object]] = []
        for start_column, column, side in BLOCK_COLUMNS:
            blocks.extend(_parse_block_column(worksheet, start_column, column, side))
        return blocks
    finally:
        workbook.close()


def save_manual_statistics(
    company_id: int,
    items: list[dict[str, object]],
    source_filename: str | None,
) -> ManualStatisticsSnapshot:
    snapshot = ManualStatisticsSnapshot.query.filter_by(company_id=company_id).first()
    if snapshot is None:
        snapshot = ManualStatisticsSnapshot(company_id=company_id)
        db.session.add(snapshot)

    snapshot.items = items
    snapshot.source_filename = (
        Path(source_filename).name[:255] if source_filename else None
    )
    db.session.flush()
    return snapshot


def manual_statistics_to_dict(
    snapshot: ManualStatisticsSnapshot | None,
) -> dict[str, object]:
    if snapshot is None:
        return {
            "items": [],
            "source_filename": None,
            "updated_at": None,
        }
    return {
        "items": list(snapshot.items or []),
        "source_filename": snapshot.source_filename,
        "updated_at": snapshot.updated_at.isoformat() if snapshot.updated_at else None,
    }

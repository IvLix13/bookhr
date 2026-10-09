from datetime import date
from io import BytesIO

import pytest
from openpyxl import load_workbook

from app.extensions import db
from app.models import Company, EmployeeGradeHistory, GradeCatalog, Reward, RewardStatus
from app.models.reward import REWARD_STATUS_LABELS
from app.services.employees import create_person_with_employment
from app.utils.dates import today_moscow


def _create_employment(company_id: int, position_grade_id: int | None = None):
    _, employment = create_person_with_employment(
        company_id=company_id,
        full_name="Иван Иванов",
        hire_date=date(2020, 1, 1),
        title="Инженер",
        position_grade_id=position_grade_id,
        education_status="yes",
    )
    db.session.commit()
    return employment


def test_create_reward_hr(hr_client, seed_company):
    employment = _create_employment(seed_company.id)

    response = hr_client.post(
        "/api/rewards",
        json={
            "employment_id": employment.id,
            "reward_type": "Благодарность",
            "status": "not_delivered",
            "directive_text": "Указ №1 от 01.01.2026",
            "notes": "Тест",
        },
    )
    assert response.status_code == 201
    payload = response.get_json()["data"]
    assert payload["reward_type"] == "Благодарность"
    assert payload["status"] == RewardStatus.NOT_DELIVERED.value
    assert payload["full_name"] == "Иван Иванов"
    assert payload["directive_text"] == "Указ №1 от 01.01.2026"


def test_create_reward_viewer_forbidden(viewer_client, seed_company):
    employment = _create_employment(seed_company.id)

    response = viewer_client.post(
        "/api/rewards",
        json={
            "employment_id": employment.id,
            "reward_type": "Благодарность",
        },
    )
    assert response.status_code == 403


def test_create_reward_accepts_extra_status(hr_client, seed_company):
    person, employment = create_person_with_employment(
        company_id=seed_company.id,
        full_name="Extra Status",
        hire_date=date(2020, 1, 1),
        title="Инженер",
    )
    db.session.commit()

    response = hr_client.post(
        "/api/rewards",
        json={
            "employment_id": employment.id,
            "reward_type": "Грамота",
            "status": RewardStatus.EXTRA_2.value,
        },
    )
    assert response.status_code == 201
    assert response.get_json()["data"]["status"] == RewardStatus.EXTRA_2.value


def test_update_reward_status_to_delivered_sets_date(hr_client, seed_company, monkeypatch):
    employment = _create_employment(seed_company.id)
    monkeypatch.setattr("app.services.rewards.today_moscow", lambda: date(2026, 7, 28))

    create_response = hr_client.post(
        "/api/rewards",
        json={
            "employment_id": employment.id,
            "reward_type": "Премия",
        },
    )
    reward_id = create_response.get_json()["data"]["id"]

    response = hr_client.patch(
        f"/api/rewards/{reward_id}",
        json={"status": RewardStatus.DELIVERED.value},
    )
    assert response.status_code == 200
    payload = response.get_json()["data"]
    assert payload["status"] == RewardStatus.DELIVERED.value
    assert payload["delivered_date"] == "2026-07-28"


def test_update_reward_preserves_delivered_date(hr_client, seed_company):
    employment = _create_employment(seed_company.id)

    create_response = hr_client.post(
        "/api/rewards",
        json={
            "employment_id": employment.id,
            "reward_type": "Премия",
            "status": RewardStatus.DELIVERED.value,
            "delivered_date": "2026-06-01",
        },
    )
    reward_id = create_response.get_json()["data"]["id"]

    response = hr_client.patch(
        f"/api/rewards/{reward_id}",
        json={"notes": "Обновлено"},
    )
    assert response.status_code == 200
    payload = response.get_json()["data"]
    assert payload["delivered_date"] == "2026-06-01"
    assert payload["notes"] == "Обновлено"


def test_list_rewards_filter_by_company(admin_client, seed_company, app):
    employment = _create_employment(seed_company.id)
    with app.app_context():
        db.session.add(
            Reward(
                employment_id=employment.id,
                reward_type="Грамота",
                status=RewardStatus.IN_HR.value,
            )
        )
        db.session.commit()

    response = admin_client.get(f"/api/rewards?company_id={seed_company.id}")
    assert response.status_code == 200
    items = response.get_json()["data"]["items"]
    assert len(items) == 1
    assert items[0]["reward_type"] == "Грамота"


def test_create_reward_requires_type(hr_client, seed_company):
    employment = _create_employment(seed_company.id)

    response = hr_client.post(
        "/api/rewards",
        json={
            "employment_id": employment.id,
            "reward_type": "   ",
        },
    )
    assert response.status_code == 400


def test_employee_list_includes_eligible_date(admin_client, seed_company, app):
    with app.app_context():
        grade_a = GradeCatalog(name="Junior", rank=1, min_years=1)
        grade_b = GradeCatalog(name="Middle", rank=2, min_years=1)
        db.session.add_all([grade_a, grade_b])
        db.session.flush()

        employment = _create_employment(seed_company.id, position_grade_id=grade_b.id)
        db.session.add(
            EmployeeGradeHistory(
                employment_id=employment.id,
                grade_id=grade_a.id,
                assigned_date=date(2025, 1, 15),
            )
        )
        db.session.commit()
        employment_id = employment.id

    response = admin_client.get(f"/api/employees?company_id={seed_company.id}&per_page=50")
    assert response.status_code == 200
    items = response.get_json()["data"]["items"]
    employee = next(item for item in items if item["id"] == employment_id)
    assert employee["eligible_date"] == "2026-01-15"


def test_reward_status_labels_cover_every_status():
    assert set(REWARD_STATUS_LABELS) == {status.value for status in RewardStatus}


def test_reward_statistics_counts_statuses_for_current_company(hr_client, seed_company):
    employment = _create_employment(seed_company.id)
    with hr_client.application.app_context():
        other = Company(name="Другая")
        db.session.add(other)
        db.session.commit()
        other_id = other.id
        other_employment = _create_employment(other_id)
        db.session.add_all([
            Reward(employment_id=employment.id, reward_type="Благодарность", status=RewardStatus.NOT_DELIVERED.value),
            Reward(employment_id=employment.id, reward_type="Грамота", status=RewardStatus.NOT_DELIVERED.value),
            Reward(employment_id=employment.id, reward_type="Медаль", status=RewardStatus.DELIVERED.value),
            Reward(employment_id=other_employment.id, reward_type="Чужая", status=RewardStatus.IN_HR.value),
        ])
        db.session.commit()

    payload = hr_client.get("/api/rewards/statistics").get_json()["data"]
    counts = {item["status"]: item["count"] for item in payload["items"]}
    assert counts["not_delivered"] == 2
    assert counts["delivered"] == 1
    assert counts["in_hr"] == 0
    assert counts["extra_1"] == 0
    assert payload["total"] == 3
    assert payload["delivered"] == 1
    assert payload["undelivered"] == 2
    assert [item["label"] for item in payload["items"]] == [
        "Не вручено",
        "В кадрах",
        "Вручено",
        "Доп. статус 1",
        "Доп. статус 2",
        "Доп. статус 3",
    ]


def test_reward_statistics_requires_login(client):
    assert client.get("/api/rewards/statistics").status_code == 401
    assert client.get("/api/rewards/statistics/export").status_code == 401


def test_reward_statistics_export_for_hr(hr_client, seed_company):
    employment = _create_employment(seed_company.id)
    db.session.add(Reward(
        employment_id=employment.id,
        reward_type="Благодарность",
        status=RewardStatus.IN_HR.value,
    ))
    db.session.commit()

    assert hr_client.get("/api/rewards/statistics").status_code == 200
    response = hr_client.get("/api/rewards/statistics/export")
    assert response.status_code == 200
    assert response.mimetype == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert "rewards_statistics.xlsx" in response.headers["Content-Disposition"]
    sheet = load_workbook(BytesIO(response.data)).active
    assert sheet.title == "Статистика поощрений"
    rows = list(sheet.values)
    assert rows[0] == ("Итоги", "Количество")
    assert rows[1:4] == [("Всего", 1), ("Вручено", 0), ("Не вручено", 1)]
    assert rows[4][0] == 'Не вручено — все поощрения, кроме состояния “Вручено”'
    assert rows[6][0] == "Разбивка по состояниям"
    assert rows[7] == ("Состояние", "Количество")
    assert ("Состояние: Не вручено", 0) in rows[8:]
    assert ("В кадрах", 1) in rows[8:]


@pytest.mark.parametrize(
    "counts, expected_total, expected_delivered",
    [
        ({}, 0, 0),
        ({"delivered": 3}, 3, 3),
        ({"in_hr": 2, "extra_2": 1}, 3, 0),
        ({"legacy_status": 2, "": 1, "delivered": 1}, 4, 1),
        ({"not_delivered": 10, "in_hr": 20, "delivered": 40,
          "extra_1": 10, "extra_2": 10, "extra_3": 10}, 100, 40),
    ],
)
def test_reward_delivery_totals_and_export(hr_client, seed_company, counts, expected_total, expected_delivered):
    employment = _create_employment(seed_company.id)
    for status, count in counts.items():
        db.session.add_all([
            Reward(
                employment_id=employment.id,
                reward_type="Грамота",
                status=status,
                # The status, not a retained delivery date, defines the totals.
                delivered_date=date(2026, 1, 1),
            )
            for _ in range(count)
        ])
    db.session.commit()

    response = hr_client.get("/api/rewards/statistics?q=ignored&per_page=1")
    assert response.status_code == 200
    stats = response.get_json()["data"]
    assert stats["total"] == expected_total
    assert stats["delivered"] == expected_delivered
    assert stats["undelivered"] == expected_total - expected_delivered
    assert stats["total"] == sum(item["count"] for item in stats["items"])
    for item in stats["items"]:
        assert item["count"] == counts.get(item["status"], 0)
    assert set(counts).issubset({item["status"] for item in stats["items"]})

    export = hr_client.get("/api/rewards/statistics/export")
    assert export.status_code == 200
    rows = list(load_workbook(BytesIO(export.data)).active.values)
    assert rows[1:4] == [
        ("Всего", stats["total"]),
        ("Вручено", stats["delivered"]),
        ("Не вручено", stats["undelivered"]),
    ]
    assert rows[8:] == [
        ("Состояние: Не вручено" if item["status"] == "not_delivered" else item["label"], item["count"])
        for item in stats["items"]
    ]


def test_reward_statistics_viewer_forbidden(viewer_client):
    assert viewer_client.get("/api/rewards/statistics").status_code == 403
    assert viewer_client.get("/api/rewards/statistics/export").status_code == 403


def test_reward_statistics_admin_allowed(admin_client):
    assert admin_client.get("/api/rewards/statistics").status_code == 200
    assert admin_client.get("/api/rewards/statistics/export").status_code == 200

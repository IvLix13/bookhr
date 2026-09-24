from datetime import date
import importlib
from io import BytesIO

import pytest
import sqlalchemy as sa
from alembic.migration import MigrationContext
from alembic.operations import Operations
from openpyxl import load_workbook

from app.extensions import db
from app.models import AuditLog, Company, Employment, Person, PersonNameHistory, GradeCatalog, EmployeeGradeHistory
from app.models.onboarding import OnboardingCell, OnboardingColumn, OnboardingPlan
from app.services.purge_employees import purge_all_employees


def employee(company_id, name="Иванов Иван"):
    person = Person()
    db.session.add(person)
    db.session.flush()
    db.session.add(PersonNameHistory(person_id=person.id, full_name=name, valid_from=date(2020, 1, 1)))
    row = Employment(person_id=person.id, company_id=company_id, hire_date=date(2020, 1, 1))
    db.session.add(row)
    db.session.commit()
    return row.id


def data(response, status=200):
    assert response.status_code == status, response.get_json()
    return response.get_json()["data"]


def setup_plan(client, company_id, kind="stage"):
    with client.application.app_context():
        employment_id = employee(company_id)
    plan = data(client.post("/api/onboarding/plans", json={"employment_id": employment_id}), 201)
    column = data(client.post("/api/onboarding/columns", json={"title": "NDA", "field_type": kind}), 201)
    return plan, column


def patch_cell(client, plan, column, **payload):
    return client.patch(f'/api/onboarding/plans/{plan["id"]}/cells/{column["id"]}',
                        json={"version": 0, "column_version": column["version"], **payload})


def test_stage_lifecycle_and_audit(hr_client, seed_company):
    plan, column = setup_plan(hr_client, seed_company.id)
    cell = data(patch_cell(hr_client, plan, column, planned_date="2026-10-01"))
    done = data(patch_cell(hr_client, plan, column, version=cell["version"], is_completed=True, completed_date="2026-09-19"))
    assert done["planned_date"] == "2026-10-01"
    assert done["is_completed"] is True
    undone = data(patch_cell(hr_client, plan, column, version=done["version"], is_completed=False))
    assert undone["completed_date"] is None
    assert undone["planned_date"] == "2026-10-01"
    result = data(hr_client.get("/api/onboarding/plans"))
    assert result["items"][0]["cells"][str(column["id"])] == undone
    assert data(hr_client.get(f'/api/onboarding/plans/{plan["id"]}'))["full_name"] == "Иванов Иван"
    with hr_client.application.app_context():
        logs = AuditLog.query.filter_by(entity_type="onboarding_cell").all()
        assert len(logs) == 3
        assert logs[-1].old_value["completed_date"] == "2026-09-19"
        assert logs[-1].user_id is not None


def test_completion_without_plan_and_date_edit(hr_client, seed_company):
    plan, column = setup_plan(hr_client, seed_company.id)
    cell = data(patch_cell(hr_client, plan, column, is_completed=True, completed_date="2026-09-19"))
    assert cell["planned_date"] is None
    changed = data(patch_cell(hr_client, plan, column, version=cell["version"], completed_date="2026-09-18"))
    assert changed["completed_date"] == "2026-09-18"


@pytest.mark.parametrize("payload", [
    {"is_completed": True}, {"planned_date": "bad"}, {"completed_date": "2026-02-30"},
    {"text_value": "wrong type"}, {"is_completed": "false"}, {"unknown": "x"},
])
def test_invalid_cells(hr_client, seed_company, payload):
    plan, column = setup_plan(hr_client, seed_company.id)
    assert patch_cell(hr_client, plan, column, **payload).status_code == 400
    with hr_client.application.app_context():
        assert OnboardingCell.query.count() == 0


@pytest.mark.parametrize("kind,values", [("text", {"text_value": "IT"}), ("date", {"date_value": "2026-09-19"})])
def test_types_and_clear(hr_client, seed_company, kind, values):
    plan, column = setup_plan(hr_client, seed_company.id, kind)
    cell = data(patch_cell(hr_client, plan, column, **values))
    key = next(iter(values))
    assert cell[key] == values[key]
    cleared = data(patch_cell(hr_client, plan, column, version=cell["version"], **{key: None}))
    assert cleared[key] is None


def test_stale_cell_and_column(hr_client, seed_company):
    plan, column = setup_plan(hr_client, seed_company.id)
    data(patch_cell(hr_client, plan, column, planned_date="2026-10-01"))
    assert patch_cell(hr_client, plan, column, planned_date="2026-10-02").status_code == 409
    updated = data(hr_client.patch(f'/api/onboarding/columns/{column["id"]}', json={"version": column["version"], "title": "New"}))
    assert updated["version"] > column["version"]
    assert hr_client.patch(f'/api/onboarding/columns/{column["id"]}', json={"version": column["version"], "title": "Lost"}).status_code == 409
    assert patch_cell(hr_client, plan, column, version=1, planned_date="2026-10-03").status_code == 409


def test_archive_restore_rename_reorder_and_type_guard(hr_client, seed_company):
    plan, column = setup_plan(hr_client, seed_company.id)
    cell = data(patch_cell(hr_client, plan, column, planned_date="2026-10-01"))
    url = f'/api/onboarding/columns/{column["id"]}'
    assert hr_client.patch(url, json={"version": 1, "field_type": "text"}).status_code == 400
    archived = data(hr_client.patch(url, json={"version": 1, "is_archived": True}))
    assert patch_cell(hr_client, plan, archived, version=cell["version"], planned_date="2026-10-02").status_code == 409
    assert hr_client.patch(url, json={"version": archived["version"], "field_type": "text"}).status_code == 400
    restored = data(hr_client.patch(url, json={"version": archived["version"], "is_archived": False, "title": "Agreement"}))
    second = data(hr_client.post("/api/onboarding/columns", json={"title": "Отдел", "field_type": "text"}), 201)
    order = [{"id": c["id"], "version": c["version"]} for c in [second, restored]]
    data(hr_client.patch("/api/onboarding/columns/order", json={"columns": order}))
    assert [c["id"] for c in data(hr_client.get("/api/onboarding/columns"))] == [second["id"], restored["id"]]
    cells = data(hr_client.get("/api/onboarding/plans"))["items"][0]["cells"]
    assert cells[str(column["id"])]["planned_date"] == "2026-10-01"
    assert str(second["id"]) not in cells
    assert hr_client.patch("/api/onboarding/columns/order", json={"columns": order}).status_code == 409
    assert hr_client.patch("/api/onboarding/columns/order", json={"columns": []}).status_code == 409


def test_empty_type_change_and_protected_fields(admin_client, seed_company):
    plan, column = setup_plan(admin_client, seed_company.id, "text")
    data(patch_cell(admin_client, plan, column, text_value=""))
    data(admin_client.patch(f'/api/onboarding/columns/{column["id"]}', json={"version": 1, "field_type": "date"}))
    assert admin_client.patch(f'/api/onboarding/plans/{plan["id"]}', json={"full_name": "hack"}).status_code == 405
    assert admin_client.post("/api/onboarding/columns", json={"title": " ", "field_type": "text"}).status_code == 400
    assert admin_client.post("/api/onboarding/columns", json={"title": "x", "field_type": "sql"}).status_code == 400


def test_tenant_isolation(hr_client, seed_company):
    plan, column = setup_plan(hr_client, seed_company.id)
    with hr_client.application.app_context():
        company = Company(name="Other")
        db.session.add(company)
        db.session.commit()
        other_employee = employee(company.id)
        other_column = OnboardingColumn(company_id=company.id, title="Secret", field_type="text", sort_order=0)
        other_plan = OnboardingPlan(employment_id=other_employee)
        db.session.add_all([other_column, other_plan])
        db.session.commit()
        cid, pid = other_column.id, other_plan.id
    assert len(data(hr_client.get("/api/onboarding/columns"))) == 1
    assert data(hr_client.get("/api/onboarding/plans"))["total"] == 1
    assert hr_client.post("/api/onboarding/plans", json={"employment_id": other_employee}).status_code == 404
    assert hr_client.get(f"/api/onboarding/plans/{pid}").status_code == 404
    assert hr_client.patch(f"/api/onboarding/columns/{cid}", json={"version": 1, "title": "hack"}).status_code == 404
    assert patch_cell(hr_client, plan, {"id": cid, "version": 1}, text_value="hack").status_code == 404
    assert patch_cell(hr_client, {"id": pid}, column, planned_date=None).status_code == 404


def test_viewer_and_anonymous_permissions(viewer_client):
    assert viewer_client.get("/api/onboarding/plans").status_code == 200
    assert viewer_client.get("/api/onboarding/columns").status_code == 200
    for method, path in [("post", "plans"), ("post", "columns"), ("patch", "columns/1"), ("patch", "columns/order"), ("patch", "plans/1/cells/1")]:
        assert getattr(viewer_client, method)(f"/api/onboarding/{path}", json={}).status_code == 403


def test_anonymous_cannot_read(client):
    assert client.get("/api/onboarding/plans").status_code == 401
    assert client.get("/api/onboarding/columns").status_code == 401
    assert client.get("/api/onboarding/export").status_code == 401


def test_viewer_can_export_valid_workbook(viewer_client):
    response = viewer_client.get("/api/onboarding/export")
    assert response.status_code == 200
    assert response.mimetype == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert "onboarding_plan.xlsx" in response.headers["Content-Disposition"]
    workbook = load_workbook(BytesIO(response.data))
    assert workbook.active.title == "План обучения"
    assert workbook.active.cell(1, 1).value == "Этап / сотрудник"


def test_export_contains_all_plans_visible_columns_and_states(hr_client, seed_company):
    plan, completed_column = setup_plan(hr_client, seed_company.id)
    data(patch_cell(
        hr_client,
        plan,
        completed_column,
        is_completed=True,
        completed_date="2026-09-19",
    ))
    text_column = data(hr_client.post(
        "/api/onboarding/columns",
        json={"title": "Отдел", "field_type": "text"},
    ), 201)
    data(patch_cell(hr_client, plan, text_column, text_value="Разработка"))
    date_column = data(hr_client.post(
        "/api/onboarding/columns",
        json={"title": "Дата", "field_type": "date"},
    ), 201)
    data(patch_cell(hr_client, plan, date_column, date_value="2026-09-20"))
    checkbox_column = data(hr_client.post(
        "/api/onboarding/columns",
        json={"title": "Ознакомление", "field_type": "checkbox"},
    ), 201)
    data(patch_cell(hr_client, plan, checkbox_column, is_completed=True))
    not_required_column = data(hr_client.post(
        "/api/onboarding/columns",
        json={"title": "Не требуется", "field_type": "checkbox"},
    ), 201)
    data(patch_cell(hr_client, plan, not_required_column, is_not_required=True))
    archived_column = data(hr_client.post(
        "/api/onboarding/columns",
        json={"title": "Архив", "field_type": "checkbox"},
    ), 201)
    data(hr_client.patch(
        f'/api/onboarding/columns/{archived_column["id"]}',
        json={"version": archived_column["version"], "is_archived": True},
    ))

    with hr_client.application.app_context():
        for index in range(25):
            employment_id = employee(seed_company.id, f"Сотрудник {index:02d}")
            db.session.add(OnboardingPlan(employment_id=employment_id))
        other_company = Company(name="Export isolation")
        db.session.add(other_company)
        db.session.commit()
        other_employment_id = employee(other_company.id, "Чужой сотрудник")
        db.session.add(OnboardingPlan(employment_id=other_employment_id))
        db.session.commit()

    response = hr_client.get("/api/onboarding/export?page=1&per_page=1&q=Несуществующий")
    assert response.status_code == 200
    sheet = load_workbook(BytesIO(response.data)).active

    assert sheet.max_column == 27
    assert sheet.cell(1, 2).value == "Иванов Иван"
    assert sheet.cell(2, 27).value == 26
    assert "Чужой сотрудник" not in [cell.value for cell in sheet[1]]
    assert [sheet.cell(row, 1).value for row in range(1, sheet.max_row + 1)] == [
        "Этап / сотрудник", "№", "Грейд", "NDA", "Отдел", "Дата",
        "Ознакомление", "Не требуется",
    ]
    assert sheet.cell(4, 2).value == "Выполнено: 19.09.2026"
    assert sheet.cell(4, 2).fill.fgColor.rgb.endswith("E2F5E9")
    assert sheet.cell(5, 2).value == "Разработка"
    assert sheet.cell(6, 2).value.strftime("%d.%m.%Y") == "20.09.2026"
    assert sheet.cell(6, 2).number_format == "DD.MM.YYYY"
    assert sheet.cell(7, 2).value == "Выполнено"
    assert sheet.cell(7, 2).fill.fgColor.rgb.endswith("E2F5E9")
    assert sheet.cell(8, 2).value == "Не нужно"


def test_duplicate_dismissal_deletion_and_rehire(hr_client, seed_company):
    plan, _ = setup_plan(hr_client, seed_company.id)
    assert hr_client.post("/api/onboarding/plans", json={"employment_id": plan["employment_id"]}).status_code == 409
    assert hr_client.delete(f'/api/employees/{plan["employment_id"]}').status_code == 400
    with hr_client.application.app_context():
        with pytest.raises(ValueError):
            purge_all_employees()
        employment = db.session.get(Employment, plan["employment_id"])
        employment.status = "dismissed"
        again = Employment(person_id=employment.person_id, company_id=seed_company.id, hire_date=date(2026, 9, 19))
        db.session.add(again)
        db.session.commit()
        new_id = again.id
    data(hr_client.post("/api/onboarding/plans", json={"employment_id": new_id}), 201)
    listed = data(hr_client.get("/api/onboarding/plans"))
    assert listed["total"] == 2
    assert listed["items"][0]["employment_status"] == "dismissed"


def test_current_name_grade_search_pagination(hr_client, seed_company):
    plan, _ = setup_plan(hr_client, seed_company.id)
    with hr_client.application.app_context():
        employment = db.session.get(Employment, plan["employment_id"])
        employment.person.name_history[0].full_name = "Петров"
        grade = GradeCatalog(name="Junior", rank=1, min_years=1)
        db.session.add(grade)
        db.session.flush()
        db.session.add(EmployeeGradeHistory(employment_id=employment.id, grade_id=grade.id, assigned_date=date(2026, 1, 1)))
        db.session.commit()
    rows = data(hr_client.get("/api/onboarding/plans?q=Петров&per_page=1"))
    assert rows["items"][0]["grade"] == "Junior"
    assert rows["items"][0]["full_name"] == "Петров"
    assert data(hr_client.get("/api/onboarding/plans?q=Несуществующий"))["total"] == 0
    assert data(hr_client.get("/api/onboarding/plans?page=2&per_page=1"))["items"] == []


def test_migration_preserves_existing_data():
    migration = importlib.import_module("migrations.versions.0016_onboarding")
    engine = sa.create_engine("sqlite://")
    with engine.begin() as conn:
        conn.exec_driver_sql("CREATE TABLE companies (id INTEGER PRIMARY KEY)")
        conn.exec_driver_sql("CREATE TABLE employments (id INTEGER PRIMARY KEY)")
        conn.exec_driver_sql("INSERT INTO companies VALUES (1)")
        conn.exec_driver_sql("INSERT INTO employments VALUES (1)")
        with Operations.context(MigrationContext.configure(conn)):
            migration.upgrade()
            assert conn.exec_driver_sql("SELECT COUNT(*) FROM onboarding_columns").scalar() == 6
            assert conn.exec_driver_sql("SELECT COUNT(*) FROM onboarding_plans").scalar() == 0
            assert conn.exec_driver_sql("SELECT COUNT(*) FROM employments").scalar() == 1
            migration.downgrade()
            migration.upgrade()
            assert conn.exec_driver_sql("SELECT COUNT(*) FROM onboarding_columns").scalar() == 6


def test_list_queries_do_not_grow_per_row_or_cell(hr_client, seed_company):
    plan, column = setup_plan(hr_client, seed_company.id)
    data(patch_cell(hr_client, plan, column, planned_date="2026-10-01"))
    with hr_client.application.app_context():
        engine = db.engine
        def count_queries():
            statements = []
            def record(conn, cursor, statement, parameters, context, executemany):
                if statement.lstrip().upper().startswith("SELECT"):
                    statements.append(statement)
            db.session.expire_all()
            sa.event.listen(engine, "before_cursor_execute", record)
            try:
                data(hr_client.get("/api/onboarding/plans"))
            finally:
                sa.event.remove(engine, "before_cursor_execute", record)
            return len(statements)
        before = count_queries()
        for i in range(12):
            eid = employee(seed_company.id, f"Сотрудник {i}")
            db.session.add(OnboardingPlan(employment_id=eid))
        db.session.commit()
        after = count_queries()
        assert after == before


def test_csrf_required_for_new_routes(hr_client):
    app = hr_client.application
    app.config["TESTING"] = False
    try:
        response = hr_client.post("/api/onboarding/columns", json={"title": "Отдел", "field_type": "text"})
        assert response.status_code == 403
        token = data(hr_client.get("/api/csrf"))["csrf_token"]
        data(hr_client.post("/api/onboarding/columns", headers={"X-CSRF-Token": token},
                            json={"title": "Отдел", "field_type": "text"}), 201)
    finally:
        app.config["TESTING"] = True


@pytest.mark.parametrize("kind", ["stage", "checkbox"])
def test_exemption_transitions_and_per_employee_scope(hr_client, seed_company, kind):
    plan, column = setup_plan(hr_client, seed_company.id, kind)
    extra = {"planned_date": "2020-01-01", "completed_date": "2026-09-22"} if kind == "stage" else {}
    done = data(patch_cell(hr_client, plan, column, is_completed=True, **extra))
    assert done["is_completed"] is True
    if kind == "checkbox":
        assert done["planned_date"] is None and done["completed_date"] is None
    exempt = data(patch_cell(hr_client, plan, column, version=done["version"], is_not_required=True))
    assert exempt["is_not_required"] is True and exempt["is_completed"] is False
    assert exempt["completed_date"] is None
    assert exempt["planned_date"] == extra.get("planned_date")
    assert data(hr_client.get(f'/api/onboarding/plans/{plan["id"]}'))["cells"][str(column["id"])]["is_not_required"] is True
    with hr_client.application.app_context():
        other_id = employee(seed_company.id, "Другой сотрудник")
    other = data(hr_client.post("/api/onboarding/plans", json={"employment_id": other_id}), 201)
    assert other["cells"] == {}
    pending = data(patch_cell(hr_client, plan, column, version=exempt["version"], is_not_required=False))
    assert pending["is_completed"] is False and pending["completed_date"] is None
    exempt = data(patch_cell(hr_client, plan, column, version=pending["version"], is_not_required=True))
    dates = {"completed_date": "2026-09-22"} if kind == "stage" else {}
    done = data(patch_cell(hr_client, plan, column, version=exempt["version"], is_completed=True, **dates))
    assert done["is_completed"] and not done["is_not_required"]
    with hr_client.application.app_context():
        logs = AuditLog.query.filter_by(entity_type="onboarding_cell").order_by(AuditLog.id).all()
        assert logs[1].old_value["is_completed"] is True
        assert logs[1].new_value["is_not_required"] is True
        assert logs[1].old_value["completed_date"] == extra.get("completed_date")


@pytest.mark.parametrize("kind,payload", [
    ("checkbox", {"planned_date": None}), ("checkbox", {"completed_date": "2026-09-22"}),
    ("checkbox", {"is_completed": True, "is_not_required": True}),
    ("stage", {"is_completed": True, "is_not_required": True, "completed_date": "2026-09-22"}),
    ("text", {"is_not_required": True}), ("date", {"is_not_required": True}),
    ("stage", {"is_not_required": "false"}),
])
def test_new_state_validation(hr_client, seed_company, kind, payload):
    plan, column = setup_plan(hr_client, seed_company.id, kind)
    assert patch_cell(hr_client, plan, column, **payload).status_code == 400


def test_exemption_blocks_type_change_and_respects_versions_archive(hr_client, seed_company):
    plan, column = setup_plan(hr_client, seed_company.id, "checkbox")
    cell = data(patch_cell(hr_client, plan, column, is_not_required=True))
    assert patch_cell(hr_client, plan, column, is_completed=True).status_code == 409
    url = f'/api/onboarding/columns/{column["id"]}'
    assert hr_client.patch(url, json={"version": 1, "field_type": "text"}).status_code == 400
    archived = data(hr_client.patch(url, json={"version": 1, "is_archived": True}))
    assert hr_client.patch(url, json={"version": archived["version"], "field_type": "stage"}).status_code == 400
    assert patch_cell(hr_client, plan, archived, version=cell["version"], clear=True).status_code == 409


@pytest.mark.parametrize("kind,values", [
    ("text", {"text_value": "Отдел"}), ("date", {"date_value": "2026-09-22"}),
    ("stage", {"is_completed": True, "planned_date": "2026-09-23", "completed_date": "2026-09-22"}),
    ("stage", {"is_not_required": True, "planned_date": "2026-09-23"}),
    ("checkbox", {"is_completed": True}), ("checkbox", {"is_not_required": True}),
])
def test_clear_cell_resets_values_keeps_identity_and_version(hr_client, seed_company, kind, values):
    plan, column = setup_plan(hr_client, seed_company.id, kind)
    original = data(patch_cell(hr_client, plan, column, **values))
    cleared = data(patch_cell(hr_client, plan, column, version=original["version"], clear=True))
    assert cleared["version"] > original["version"]
    assert all(cleared[key] is None for key in ["text_value", "date_value", "planned_date", "completed_date"])
    assert not cleared["is_completed"] and not cleared["is_not_required"]
    assert patch_cell(hr_client, plan, column, version=original["version"], **values).status_code == 409
    assert patch_cell(hr_client, plan, column, **values).status_code == 409
    again = data(patch_cell(hr_client, plan, column, version=cleared["version"], clear=True))
    assert again["version"] > cleared["version"]
    with hr_client.application.app_context():
        assert OnboardingPlan.query.count() == 1 and OnboardingColumn.query.count() == 1
        assert OnboardingCell.query.count() == 1
        log = AuditLog.query.filter_by(action="clear").order_by(AuditLog.id).first()
        for key, value in values.items():
            assert log.old_value[key] == value
    restored = data(patch_cell(hr_client, plan, column, version=again["version"], **values))
    for key, value in values.items():
        assert restored[key] == value


def test_clear_empty_cell_and_invalid_mixed_request(hr_client, seed_company):
    plan, column = setup_plan(hr_client, seed_company.id)
    assert patch_cell(hr_client, plan, column, clear=True, is_completed=True).status_code == 400
    cleared = data(patch_cell(hr_client, plan, column, clear=True))
    assert cleared["version"] > 0
    assert patch_cell(hr_client, plan, column, planned_date="2026-09-22").status_code == 409


def test_states_migration_preserves_values_and_guards_downgrade():
    original = importlib.import_module("migrations.versions.0016_onboarding")
    migration = importlib.import_module("migrations.versions.0017_onboarding_states")
    engine = sa.create_engine("sqlite://")
    with engine.begin() as conn:
        conn.exec_driver_sql("CREATE TABLE companies (id INTEGER PRIMARY KEY)")
        conn.exec_driver_sql("CREATE TABLE employments (id INTEGER PRIMARY KEY)")
        conn.exec_driver_sql("INSERT INTO companies VALUES (1)")
        conn.exec_driver_sql("INSERT INTO employments VALUES (1)")
        with Operations.context(MigrationContext.configure(conn)):
            original.upgrade()
            conn.exec_driver_sql("INSERT INTO onboarding_plans (id, employment_id, created_at, updated_at) VALUES (1,1,CURRENT_TIMESTAMP,CURRENT_TIMESTAMP)")
            conn.exec_driver_sql("INSERT INTO onboarding_cells (plan_id, column_id, planned_date, is_completed, completed_date, version, created_at, updated_at) VALUES (1,3,'2026-09-20',1,'2026-09-22',4,CURRENT_TIMESTAMP,CURRENT_TIMESTAMP)")
            before = conn.exec_driver_sql("SELECT planned_date, is_completed, completed_date, version FROM onboarding_cells").one()
            migration.upgrade()
            assert conn.exec_driver_sql("SELECT planned_date, is_completed, completed_date, version FROM onboarding_cells").one() == before
            assert conn.exec_driver_sql("SELECT is_not_required FROM onboarding_cells").scalar() == 0
            migration.downgrade()
            migration.upgrade()
            conn.exec_driver_sql("UPDATE onboarding_columns SET field_type='checkbox' WHERE id=4")
            with pytest.raises(RuntimeError, match="Откат запрещён"):
                migration.downgrade()
            conn.exec_driver_sql("UPDATE onboarding_columns SET field_type='stage' WHERE id=4")
            conn.exec_driver_sql("UPDATE onboarding_cells SET is_completed=0, completed_date=NULL, is_not_required=1")
            with pytest.raises(RuntimeError, match="Откат запрещён"):
                migration.downgrade()
            with pytest.raises(sa.exc.IntegrityError):
                conn.exec_driver_sql("UPDATE onboarding_cells SET is_completed=1 WHERE is_not_required=1")

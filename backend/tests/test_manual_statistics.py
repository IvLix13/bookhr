from datetime import date
from io import BytesIO

from openpyxl import Workbook, load_workbook

from app.extensions import db
from app.models import ManualStatisticsSnapshot


def _manual_statistics_file(left_blocks=(), right_blocks=()) -> BytesIO:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Ручная статистика"
    for start_column, blocks in ((1, left_blocks), (4, right_blocks)):
        row = 1
        for block_number, (title, items) in enumerate(blocks, start=1):
            worksheet.cell(row=row, column=start_column, value=f"Блок {block_number}")
            worksheet.cell(row=row + 1, column=start_column, value=title)
            worksheet.cell(row=row + 2, column=start_column, value="Наименование параметра")
            worksheet.cell(row=row + 2, column=start_column + 1, value="Значение")
            for item_number, (label, value) in enumerate(items, start=row + 3):
                worksheet.cell(row=item_number, column=start_column, value=label)
                worksheet.cell(row=item_number, column=start_column + 1, value=value)
            row += len(items) + 5
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


def _upload(client, workbook, filename="manual.xlsx"):
    return client.post(
        "/api/stats/manual/upload",
        data={"file": (workbook, filename)},
        content_type="multipart/form-data",
    )


def test_manual_stats_empty_for_viewer(viewer_client):
    response = viewer_client.get("/api/stats/manual")
    assert response.status_code == 200
    assert response.get_json()["data"] == {
        "items": [], "source_filename": None, "updated_at": None,
    }


def test_manual_stats_template_has_left_and_right_blocks(viewer_client):
    response = viewer_client.get("/api/stats/manual/template")
    assert response.status_code == 200
    workbook = load_workbook(BytesIO(response.data), read_only=True, data_only=True)
    worksheet = workbook.active
    assert (worksheet["A1"].value, worksheet["A2"].value) == ("Блок 1", "Договоры")
    assert (worksheet["A3"].value, worksheet["B3"].value) == ("Наименование параметра", "Значение")
    assert worksheet["C1"].value is None
    assert (worksheet["D1"].value, worksheet["D2"].value) == ("Блок 1", "Аттестация")
    assert (worksheet["D3"].value, worksheet["E3"].value) == ("Наименование параметра", "Значение")
    workbook.close()


def test_manual_stats_upload_parses_both_columns_and_preserves_order(hr_client, seed_company):
    response = _upload(hr_client, _manual_statistics_file(
        left_blocks=(("Договоры", (("Заключено за месяц", 15), ("Дата отчёта", date(2026, 9, 14)))),
                     ("Грейды", (("Назначено за месяц", 8),))),
        right_blocks=(("Аттестация", (("Заключено за месяц", 21), ("Комментарий", "Готово"))),),
    ))
    assert response.status_code == 200
    data = response.get_json()["data"]
    assert data["items"] == [
        {"column": "left", "title": "Договоры", "items": [
            {"label": "Заключено за месяц", "value": "15"},
            {"label": "Дата отчёта", "value": "2026-09-14T00:00:00"},
        ]},
        {"column": "left", "title": "Грейды", "items": [
            {"label": "Назначено за месяц", "value": "8"},
        ]},
        {"column": "right", "title": "Аттестация", "items": [
            {"label": "Заключено за месяц", "value": "21"},
            {"label": "Комментарий", "value": "Готово"},
        ]},
    ]
    assert data["source_filename"] == "manual.xlsx"
    assert data["updated_at"] is not None
    assert ManualStatisticsSnapshot.query.filter_by(company_id=seed_company.id).one().items == data["items"]


def test_manual_stats_upload_allows_empty_rows_and_skips_empty_blocks(hr_client):
    response = _upload(hr_client, _manual_statistics_file(
        left_blocks=(("Пустой", ()), ("Заполненный", (("Показатель", 1),))),
    ))
    assert response.status_code == 200
    assert response.get_json()["data"]["items"] == [
        {"column": "left", "title": "Заполненный", "items": [
            {"label": "Показатель", "value": "1"},
        ]},
    ]


def test_manual_stats_upload_replaces_previous_snapshot(hr_client):
    first = _upload(hr_client, _manual_statistics_file(
        left_blocks=(("Первый", (("Значение", 1),)),),
    ), "first.xlsx")
    assert first.status_code == 200
    second = _upload(hr_client, _manual_statistics_file(
        right_blocks=(("Второй", (("Значение", 2),)),),
    ), "second.xlsx")
    assert second.status_code == 200
    data = second.get_json()["data"]
    assert data["items"] == [{"column": "right", "title": "Второй", "items": [
        {"label": "Значение", "value": "2"},
    ]}]
    assert data["source_filename"] == "second.xlsx"
    assert ManualStatisticsSnapshot.query.count() == 1


def test_manual_stats_upload_rejects_missing_block_title(hr_client):
    workbook = Workbook()
    workbook.active["A1"] = "Блок 1"
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    response = _upload(hr_client, output)
    assert response.status_code == 400
    assert "левой колонке, строка 2" in response.get_json()["message"]
    assert "название блока" in response.get_json()["message"]


def test_manual_stats_upload_rejects_wrong_block_headers(hr_client):
    workbook = load_workbook(_manual_statistics_file(left_blocks=(("Блок", (("Показатель", 1),)),)))
    workbook.active["A3"] = "Не тот заголовок"
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    response = _upload(hr_client, output)
    assert response.status_code == 400
    assert "левой колонке, строка 3" in response.get_json()["message"]
    assert "ожидаются заголовки" in response.get_json()["message"]


def test_manual_stats_upload_rejects_value_without_label(hr_client):
    response = _upload(hr_client, _manual_statistics_file(
        left_blocks=(("Блок", ((None, 10),)),),
    ))
    assert response.status_code == 400
    assert "левой колонке, строка 4" in response.get_json()["message"]
    assert "не указано наименование параметра" in response.get_json()["message"]


def test_old_flat_snapshot_is_returned_without_server_error(app, viewer_client, seed_company):
    with app.app_context():
        db.session.add(ManualStatisticsSnapshot(
            company_id=seed_company.id,
            items=[{"label": "Старый показатель", "value": "10"}],
            source_filename="old.xlsx",
        ))
        db.session.commit()
    response = viewer_client.get("/api/stats/manual")
    assert response.status_code == 200
    assert response.get_json()["data"]["items"] == [
        {"label": "Старый показатель", "value": "10"},
    ]


def test_manual_stats_upload_viewer_forbidden(viewer_client):
    response = _upload(viewer_client, _manual_statistics_file(
        left_blocks=(("Блок", (("Показатель", 1),)),),
    ))
    assert response.status_code == 403
